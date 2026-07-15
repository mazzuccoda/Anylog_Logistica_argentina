#!/usr/bin/env python3
"""Genera la plantilla Excel de parametros para el modelo de logistica de jugo."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

wb = Workbook()

hdr = Font(bold=True, color="FFFFFF")
fill = PatternFill("solid", fgColor="2E5A87")
note = Font(italic=True, color="808080", size=9)
thin = Side(style="thin", color="BBBBBB")
border = Border(left=thin, right=thin, top=thin, bottom=thin)


def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = hdr
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border


def frame(ws, r1, r2, c1, c2):
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            ws.cell(row=r, column=c).border = border


# ---------- Hoja 1: Depositos ----------
ws1 = wb.active
ws1.title = "Depositos"
cols = ["id", "nombre", "region", "tarifa_almacenamiento_$tn_dia",
        "capacidad_maxima_tn", "costo_transporte_unitario_$tn",
        "tiempo_viaje_horas"]
ws1.append(cols)
style_header(ws1, 1, len(cols))
rows = [
    [0, "Tucuman_A", "Tucuman", "[COMPLETAR]", "[COMPLETAR]", "[X]", "[COMPLETAR]"],
    [1, "Tucuman_B", "Tucuman", "[COMPLETAR]", "[COMPLETAR]", "[X]", "[COMPLETAR]"],
    [2, "Tucuman_C", "Tucuman", "[COMPLETAR]", "[COMPLETAR]", "[X]", "[COMPLETAR]"],
    [3, "BsAs_A", "BuenosAires", "[COMPLETAR]", "[COMPLETAR]", "[Y]", "[COMPLETAR]"],
    [4, "BsAs_B", "BuenosAires", "[COMPLETAR]", "[COMPLETAR]", "[Y]", "[COMPLETAR]"],
]
for r in rows:
    ws1.append(r)
frame(ws1, 1, 1 + len(rows), 1, len(cols))
widths = [6, 14, 14, 30, 22, 30, 20]
for i, w in enumerate(widths, 1):
    ws1.column_dimensions[chr(64 + i)].width = w
ws1.cell(row=8, column=1, value="Nota: 3 depositos en Tucuman (costo transporte X) y 2 en Buenos Aires (costo Y).").font = note

# ---------- Hoja 2: Produccion mensual (estacionalidad) ----------
ws2 = wb.create_sheet("ProduccionMensual")
cols2 = ["mes", "nombre_mes", "produccion_diaria_tn"]
ws2.append(cols2)
style_header(ws2, 1, len(cols2))
meses = ["Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio", "Julio",
         "Agosto", "Septiembre", "Octubre", "Noviembre", "Diciembre"]
for i, m in enumerate(meses, 1):
    ws2.append([i, m, "[COMPLETAR]"])
frame(ws2, 1, 13, 1, 3)
for i, w in enumerate([6, 16, 24], 1):
    ws2.column_dimensions[chr(64 + i)].width = w
ws2.cell(row=15, column=1, value="Perfil estacional: tn/dia promedio para cada mes.").font = note

# ---------- Hoja 3: Parametros globales ----------
ws3 = wb.create_sheet("Globales")
cols3 = ["parametro", "valor", "unidad", "descripcion"]
ws3.append(cols3)
style_header(ws3, 1, len(cols3))
glob = [
    ["capacidad_camara_propia", 5000, "tn", "Capacidad maxima camara propia"],
    ["produccion_diaria_promedio", "[COMPLETAR]", "tn/dia", "Usada si no se usa perfil mensual"],
    ["costo_transporte_tucuman_X", "[COMPLETAR]", "$/tn", "Costo transporte planta->Tucuman"],
    ["costo_transporte_bsas_Y", "[COMPLETAR]", "$/tn", "Costo transporte planta->Buenos Aires"],
    ["capacidad_camion", "[COMPLETAR]", "tn", "Carga por viaje de camion"],
    ["tiempo_viaje_tucuman", "[COMPLETAR]", "horas", "Tiempo de viaje a Tucuman"],
    ["tiempo_viaje_bsas", "[COMPLETAR]", "horas", "Tiempo de viaje a Buenos Aires"],
    ["horizonte_dias", 365, "dias", "Horizonte de simulacion (1 anio)"],
]
for g in glob:
    ws3.append(g)
frame(ws3, 1, 1 + len(glob), 1, 4)
for i, w in enumerate([30, 14, 10, 45], 1):
    ws3.column_dimensions[chr(64 + i)].width = w

wb.save("/home/ubuntu/anylogic_jugo/parametros_jugo.xlsx")
print("Excel generado: parametros_jugo.xlsx")
